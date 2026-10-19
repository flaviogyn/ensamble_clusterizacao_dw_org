from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
import os

IMG = "imagens_apresentacao"
BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
CYAN = RGBColor(0x00, 0xD4, 0xFF)
PURPLE = RGBColor(0xA8, 0x55, 0xF7)
GREEN = RGBColor(0x10, 0xB9, 0x81)
YELLOW = RGBColor(0xF1, 0xC4, 0x0F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x94, 0xA3, 0xB8)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    if line:
        r.line.color.rgb = line
        r.line.width = Pt(1.25)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, t, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font="Segoe UI"):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def base(title, subtitle=None, n=None):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    rect(s, 0, 0, 0.18, 7.5, CYAN)
    rect(s, 0.18, 0, 0.06, 7.5, PURPLE)
    text(s, 0.7, 0.35, 11.5, 0.8, title, 32, WHITE, True, font="Segoe UI Semibold")
    if subtitle:
        text(s, 0.7, 1.1, 11.5, 0.5, subtitle, 16, CYAN)
    rect(s, 0.75, 1.65, 1.2, 0.05, PURPLE)
    if n:
        text(s, 12.0, 6.95, 1, 0.4, f"{n}", 12, GRAY, align=PP_ALIGN.RIGHT)
    text(s, 0.7, 6.95, 8, 0.4, "DW Organizacional · dm_vendas · K-Means + Apriori", 11, GRAY)
    return s


def bullets(s, x, y, w, h, items, size=17):
    rect(s, x, y, w, h, CARD, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
    tb = s.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.2), Inches(w - 0.5), Inches(h - 0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        head, _, rest = it.partition("::")
        r = p.add_run()
        r.text = "▸ "
        r.font.color.rgb = CYAN
        r.font.size = Pt(size)
        if rest:
            r2 = p.add_run()
            r2.text = head.strip() + " "
            r2.font.bold = True
            r2.font.color.rgb = CYAN
            r2.font.size = Pt(size)
            r2.font.name = "Segoe UI"
            r3 = p.add_run()
            r3.text = rest.strip()
            r3.font.color.rgb = WHITE
            r3.font.size = Pt(size)
            r3.font.name = "Segoe UI"
        else:
            r3 = p.add_run()
            r3.text = head
            r3.font.color.rgb = WHITE
            r3.font.size = Pt(size)
            r3.font.name = "Segoe UI"


def picture(s, fname, x, y, maxw, maxh):
    path = os.path.join(IMG, fname)
    w, h = Image.open(path).size
    ratio = min(maxw / w, maxh / h)
    pw, ph = w * ratio, h * ratio
    px = x + (maxw - pw) / 2
    py = y + (maxh - ph) / 2
    rect(s, px - 0.06, py - 0.06, pw + 0.12, ph + 0.12, CARD, MSO_SHAPE.ROUNDED_RECTANGLE, CYAN)
    s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))


def kpi(s, x, y, w, value, label, color):
    rect(s, x, y, w, 1.25, CARD, MSO_SHAPE.ROUNDED_RECTANGLE, color)
    text(s, x, y + 0.1, w, 0.7, value, 32, color, True, PP_ALIGN.CENTER)
    text(s, x, y + 0.8, w, 0.4, label, 12, GRAY, False, PP_ALIGN.CENTER)


def find(prefix):
    return [f for f in os.listdir(IMG) if f.startswith(prefix)][0]


# 1 - Capa
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
rect(s, 0, 0, 0.35, 7.5, CYAN)
rect(s, 0.35, 0, 0.12, 7.5, PURPLE)
rect(s, 9.2, -1.5, 6, 6, RGBColor(0x1E, 0x1B, 0x4B), MSO_SHAPE.OVAL)
rect(s, 10.4, 3.8, 4, 4, RGBColor(0x0B, 0x3B, 0x52), MSO_SHAPE.OVAL)
text(s, 1.1, 1.6, 9, 0.5, "MINERAÇÃO DE DADOS · DATA WAREHOUSE", 16, CYAN, True)
text(s, 1.1, 2.2, 10, 1.8, ["Clusterização de Clientes", "& Regras de Associação"], 48, WHITE, True)
rect(s, 1.15, 4.2, 1.8, 0.07, PURPLE)
text(s, 1.1, 4.4, 10, 0.6, "DW Organizacional — Mercearia (B2C) + Northwind (B2B)", 20, GRAY)
text(s, 1.1, 5.1, 10, 0.5, "K-Means + Apriori  ·  Data Mart dm_vendas  ·  Outubro/2026", 16, CYAN)

# 2 - Contexto
s = base("Contexto e Objetivos", "Duas fontes transacionais, uma visão integrada", 2)
kpi(s, 0.75, 1.95, 2.8, "100", "Clientes Mercearia (B2C)", CYAN)
kpi(s, 3.75, 1.95, 2.8, "89", "Clientes Northwind (B2B)", PURPLE)
kpi(s, 6.75, 1.95, 2.8, "189", "Base integrada via DW", GREEN)
kpi(s, 9.75, 1.95, 2.8, "930", "Cestas analisadas", YELLOW)
bullets(s, 0.75, 3.45, 11.8, 3.3, [
    "1. Clusterizar:: clientes por comportamento de compra (K-Means)",
    "2. Segmentar:: as transações por grupo de clientes",
    "3. Minerar:: Apriori aplicado dentro de cada cluster (Market Basket)",
    "4. Gerar:: regras de associação específicas por segmento",
    "Diferencial:: o Apriori NÃO roda na base toda — roda por segmento, gerando regras mais relevantes",
], 18)

# 3 - Arquitetura
s = base("Arquitetura da Solução", "Do DW às 3.804 regras de associação", 3)
picture(s, find("fluxo_solucao"), 0.75, 1.9, 7.4, 4.9)
bullets(s, 8.4, 1.9, 4.2, 4.9, [
    "Fontes:: Mercearia + Northwind → dw_organizacional",
    "Data Mart:: dm_vendas, 7 variáveis comportamentais",
    "K-Means:: K = 4 clusters",
    "Apriori:: Supp ≥ 5% · Conf ≥ 30% · Lift ≥ 1",
    "id_transacao:: sistema_origem + '_' + id_pedido (unicidade global)",
], 15)

# 4 - Preparação
s = base("Preparação e Normalização", "7 variáveis comportamentais por cliente", 4)
vars_ = [("recencia_dias", "Dias desde a última compra"), ("frequencia_compras", "Nº de transações"),
         ("quantidade_itens", "Total de itens comprados"), ("valor_total", "Valor financeiro total (R$)"),
         ("ticket_medio", "Valor médio por transação"), ("produtos_distintos", "Variedade de produtos"),
         ("categorias_distintas", "Variedade de categorias")]
for i, (v, d) in enumerate(vars_):
    col, row = i % 2, i // 2
    x, y = 0.75 + col * 3.5, 1.95 + row * 1.2
    rect(s, x, y, 3.3, 1.05, CARD, MSO_SHAPE.ROUNDED_RECTANGLE, PURPLE if i % 2 else CYAN)
    text(s, x + 0.1, y + 0.08, 3.1, 0.4, v, 15, CYAN if i % 2 == 0 else PURPLE, True, font="Consolas")
    text(s, x + 0.1, y + 0.5, 3.1, 0.5, d, 13, GRAY)
bullets(s, 7.95, 1.95, 4.65, 4.75, [
    "Por que normalizar?:: K-Means usa distância euclidiana",
    "valor_total (milhares) dominaria a frequência (dezenas)",
    "StandardScaler:: média 0 e desvio-padrão 1",
    "Resultado:: todas as variáveis com peso equivalente",
], 16)

# 5 - K-Means
s = base("K-Means e Seleção do K Ótimo", "Silhouette Score testado de K=2 a K=7", 5)
picture(s, find("kmeans_processo"), 0.75, 1.9, 7.6, 4.9)
bullets(s, 8.6, 1.9, 4.0, 4.9, [
    "Algoritmo:: atribui clientes ao centroide mais próximo e recalcula até convergir",
    "Silhouette:: mede coesão interna e separação externa",
    "Resultado:: K = 4 maximizou o score",
], 15)

# 6 - Distribuição
s = base("Distribuição dos Clusters", "189 clientes em 4 perfis distintos", 6)
picture(s, find("distribuicao_clusters"), 0.75, 1.9, 6.6, 4.9)
rows = [("0", "Mercearia", "100", "B2C · cestas de 10 itens", CYAN),
        ("1", "Northwind", "39", "B2B · muitas transações, cestas enxutas", PURPLE),
        ("2", "Northwind", "3", "B2B · clientes VIP", YELLOW),
        ("3", "Northwind", "47", "B2B · perfil intermediário", GREEN)]
for i, (c, o, n, d, col) in enumerate(rows):
    y = 1.9 + i * 1.25
    rect(s, 7.6, y, 5.0, 1.1, CARD, MSO_SHAPE.ROUNDED_RECTANGLE, col)
    text(s, 7.7, y + 0.1, 1.0, 0.9, f"C{c}", 28, col, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.7, y + 0.1, 3.8, 0.4, f"{o} · {n} clientes", 16, WHITE, True)
    text(s, 8.7, y + 0.55, 3.8, 0.5, d, 12, GRAY)

# 7 - Apriori
s = base("Pipeline Apriori", "Market Basket Analysis por cluster", 7)
steps = [("Transações\npor cluster", CYAN), ("Matriz\nbinária", PURPLE), ("Itemsets\nfrequentes", GREEN), ("Regras\nA → B", YELLOW)]
for i, (t_, col) in enumerate(steps):
    x = 0.75 + i * 1.78
    rect(s, x, 2.3, 1.55, 1.9, CARD, MSO_SHAPE.ROUNDED_RECTANGLE, col)
    text(s, x, 2.3, 1.55, 1.9, t_.split("\n"), 14, col, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 3:
        text(s, x + 1.5, 2.9, 0.4, 0.6, "›", 28, GRAY, True, PP_ALIGN.CENTER)
kpi(s, 0.75, 5.3, 2.2, "≥ 5%", "MIN_SUPPORT", CYAN)
kpi(s, 3.15, 5.3, 2.2, "≥ 30%", "MIN_CONFIDENCE", PURPLE)
kpi(s, 5.55, 5.3, 2.2, "≥ 1,0", "MIN_LIFT", GREEN)
bullets(s, 8.0, 1.9, 4.6, 4.65, [
    "Matriz:: TransactionEncoder (transações × produtos)",
    "Cluster 0:: 3.804 regras ✓",
    "Clusters 1, 2, 3:: 0 regras (suporte < 5%)",
    "Biblioteca:: MLxtend",
], 16)

# 8 - Métricas
s = base("Métricas de Associação", "Suporte · Confiança · Lift", 8)
picture(s, find("metricas_associacao"), 0.75, 1.9, 7.2, 4.9)
kpi(s, 8.3, 1.9, 4.3, "Suporte 16%", "16 de cada 100 cestas", YELLOW)
kpi(s, 8.3, 3.35, 4.3, "Confiança 100%", "sempre juntos", GREEN)
kpi(s, 8.3, 4.8, 4.3, "Lift 5,0", "5× mais provável que o acaso", PURPLE)

# 9 - Exemplo de regra
s = base("Exemplo de Regra Descoberta", "Cluster 0 — Mercearia", 9)
picture(s, find("market_basket_exemplo"), 0.75, 1.9, 6.0, 4.9)
bullets(s, 7.0, 1.9, 5.6, 4.9, [
    "Regra:: Achocolatado 400g + Água Sanitária 1L → Água Mineral 1.5L",
    "Açúcar + Desinfetante:: → Creme Dental",
    "Amaciante + Feijão:: → Condicionador",
    "Arroz + Sabão em Pó:: → Shampoo",
    "Top 20 regras:: Suporte 16% · Confiança 100% · Lift 5,0",
], 15)

# 10 - Comparativo
s = base("Comparativo B2C vs B2B", "Por que o B2B não gerou regras com suporte de 5%?", 10)
picture(s, find("comparativo_clusters"), 0.75, 1.9, 6.4, 4.9)
bullets(s, 7.4, 1.9, 5.2, 2.7, [
    "Cestas B2B:: 2–3 itens vs 10 no B2C",
    "Suporte máx.:: 7,1% (B2B) vs 30% (B2C)",
    "Catálogo:: disperso e gourmet",
], 15)
bullets(s, 7.4, 4.8, 5.2, 1.95, [
    "Solução:: reduzir MIN_SUPPORT para 1–2%",
    "Ex.:: Chai → Flotemysost",
], 15)

# 11 - Conclusões
s = base("Aplicações e Conclusões", "De relatórios históricos à inteligência prescritiva", 11)
apps = [("🛍 Promoções Cruzadas", CYAN), ("🏪 Layout de Gôndola", PURPLE),
        ("🤖 Recomendação no Checkout", GREEN), ("📦 Kits Promocionais", YELLOW)]
for i, (t_, col) in enumerate(apps):
    x = 0.75 + i * 3.0
    rect(s, x, 1.95, 2.8, 1.2, CARD, MSO_SHAPE.ROUNDED_RECTANGLE, col)
    text(s, x, 1.95, 2.8, 1.2, t_, 16, col, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
bullets(s, 0.75, 3.4, 11.8, 3.35, [
    "Resultado:: 189 clientes → 4 clusters → 930 cestas → 3.804 regras no PostgreSQL",
    "Lição 1:: a granularidade do id_transacao deve ser o pedido, não o item",
    "Lição 2:: B2C e B2B exigem calibrações distintas de MIN_SUPPORT",
    "Lição 3:: o Lift garante associações genuínas, não apenas produtos populares",
    "Conclusão:: o dm_vendas agora alimenta motores de recomendação e decisões de marketing",
], 16)

prs.save("Apresentacao_Moderna_Clusterizacao_Apriori.pptx")
print("OK")
